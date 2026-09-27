"""Render Tugas 4 charts and a DOCX companion from its Markdown report."""

from __future__ import annotations

import csv
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "Laporan"
ASSET_DIR = REPORT_DIR / "Tugas_4_Assets"
MD_PATH = REPORT_DIR / "Tugas_4_Laporan_Akhir.md"
DOCX_PATH = REPORT_DIR / "Tugas_4_Laporan_Akhir.docx"

NAVY = "15324B"
TEAL = "158C8C"
GOLD = "E3A72F"
RED = "C44E52"
GREEN = "3A9265"
GRAY = "5D6B78"


def theme() -> None:
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9,
        "axes.titlesize": 11, "axes.labelcolor": f"#{GRAY}",
        "text.color": f"#{NAVY}", "axes.edgecolor": "#D7E0E7",
        "axes.spines.top": False, "axes.spines.right": False,
        "grid.color": "#E8EDF1", "figure.facecolor": "white",
        "axes.facecolor": "white", "savefig.facecolor": "white",
    })


def architecture_figure() -> None:
    fig, ax = plt.subplots(figsize=(12, 5.3))
    ax.set_xlim(0, 12); ax.set_ylim(0, 5.3); ax.axis("off")
    stages = [
        (0.3, 3.45, 1.7, 1.0, "Gambar", "foto / frame", NAVY),
        (2.55, 3.45, 1.8, 1.0, "Deteksi ikan", "YOLO26n + DAM", TEAL),
        (4.9, 3.45, 1.7, 1.0, "Crop ikan", "satu crop / box", NAVY),
        (7.15, 3.95, 1.8, 0.9, "Deteksi lesi", "YOLO26n + DAM", RED),
        (7.15, 2.85, 1.8, 0.9, "Klasifikasi", "YOLO26n-cls + DAM", GOLD),
        (9.7, 3.45, 1.9, 1.0, "Dashboard", "prediksi untuk tinjauan", NAVY),
    ]
    for x, y, w, h, title, sub, color in stages:
        ax.add_patch(plt.Rectangle((x, y), w, h, facecolor="white", edgecolor=f"#{color}", linewidth=2, joinstyle="round"))
        ax.text(x+w/2, y+h*.64, title, ha="center", va="center", weight="bold", color=f"#{color}", fontsize=10)
        ax.text(x+w/2, y+h*.3, sub, ha="center", va="center", color=f"#{GRAY}", fontsize=8)
    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2,y2), xytext=(x1,y1), arrowprops={"arrowstyle":"-|>","lw":1.8,"color":f"#{GRAY}"})
    arrow(2.0,3.95,2.5,3.95); arrow(4.4,3.95,4.85,3.95)
    arrow(6.65,3.95,7.1,4.35); arrow(6.65,3.95,7.1,3.3)
    arrow(8.98,4.35,9.65,4.05); arrow(8.98,3.3,9.65,3.85)
    ax.text(6.0, 5.12, "Tahap berantai: kegagalan deteksi ikan dapat menghentikan analisis crop", ha="center", fontsize=9, color=f"#{RED}")

    ax.add_patch(plt.Rectangle((0.45, 0.45), 11.05, 1.55, facecolor="#F3F7FA", edgecolor="#C8D6E0", linewidth=1.2))
    ax.text(.7, 1.67, "Graph deteksi proyek", color=f"#{NAVY}", weight="bold", fontsize=10)
    ax.text(.7, 1.32, "Backbone: Conv → C3k2 → SPPF → C2PSA", color=f"#{GRAY}", fontsize=8.5)
    ax.text(.7, .88, "DAM pada fitur terdalam; Detect memakai fitur multi-skala.", color=f"#{GRAY}", fontsize=8.5)
    ax.text(6.4, 1.67, "Graph klasifikasi proyek", color=f"#{NAVY}", weight="bold", fontsize=10)
    ax.text(6.4, 1.32, "Backbone: Conv → C3k2 → C2PSA → DAM → Classify", color=f"#{GRAY}", fontsize=8.5)
    ax.text(6.4, .88, "Keluaran kelas crop; bukan lokasi lesi/diagnosis lab.", color=f"#{GRAY}", fontsize=8.5)
    ax.text(6, .15, "Confidence / skor model tidak sama dengan akurasi atau probabilitas klinis.", ha="center", color=f"#{RED}", fontsize=9, weight="bold")
    fig.tight_layout()
    fig.savefig(ASSET_DIR / "arsitektur_pipeline.png", dpi=200, bbox_inches="tight")
    plt.close(fig)


def metric_figure() -> None:
    fig, axes = plt.subplots(1, 3, figsize=(14, 4.6), gridspec_kw={"width_ratios":[1.15,1.15,.85]})
    labels = ["YOLO11n\ntrained", "YOLO26n\nzero-shot", "YOLO26n +\nDAM trained"]
    colors = [f"#{NAVY}", f"#{GOLD}", f"#{TEAL}"]
    metrics = {
        "Deteksi ikan": [[.3592,.9996,.3618,.3379],[.1971,.3551,.1531,.1024],[.3601,.9968,.3757,.3446]],
        "Deteksi lesi": [[.6733,.5646,.6263,.2968],[.0574,.0522,.0041,.0017],[.5967,.4528,.4899,.1913]],
    }
    names=["Precision","Recall","mAP50","mAP50–95"]
    for ax,(title,rows) in zip(axes[:2],metrics.items()):
        vals=np.asarray(rows)
        x=np.arange(4); width=.23
        for i,(lab,color) in enumerate(zip(labels,colors)):
            bars=ax.bar(x+(i-1)*width,vals[i],width,label=lab.replace("\n"," "),color=color)
            for b,v in zip(bars,vals[i]):
                if v>=.09: ax.text(b.get_x()+b.get_width()/2,v+.025,f"{v:.2f}",ha="center",va="bottom",fontsize=7,rotation=0)
        ax.set_title(title,weight="bold"); ax.set_xticks(x,names,rotation=18,ha="right")
        ax.set_ylim(0,1.16); ax.set_ylabel("Validation score"); ax.grid(axis="y",alpha=.8); ax.set_axisbelow(True)
    cls=[.9929,.0357,.9943]
    ax=axes[2]
    bars=ax.bar(np.arange(3),cls,color=colors,width=.65)
    ax.set_ylim(0,1.13); ax.set_xticks(np.arange(3),labels,rotation=15,ha="right")
    ax.set_title("Klasifikasi · Top-1",weight="bold"); ax.set_ylabel("Top-1 validation"); ax.grid(axis="y",alpha=.8); ax.set_axisbelow(True)
    for b,v in zip(bars,cls): ax.text(b.get_x()+b.get_width()/2,v+.025,f"{v:.3f}",ha="center",fontsize=8)
    axes[2].set_xticks(np.arange(3),["YOLO11n-cls\ntrained","YOLO26n-cls\nzero-shot","YOLO26n-cls +\nDAM trained"],rotation=0,ha="center",fontsize=7.5)
    handles,leglabels=axes[0].get_legend_handles_labels()
    fig.legend(handles,leglabels,loc="upper center",bbox_to_anchor=(.5,1.02),ncol=3,frameon=False)
    fig.suptitle("Perbandingan skor yang tersedia — protokol eksperimen berbeda",y=1.12,fontsize=14,weight="bold",color=f"#{NAVY}")
    fig.text(.5,.055,"Klasifikasi: overlap validation/train dilaporkan; top-1 tinggi bukan estimasi independen.",ha="center",fontsize=8,color=f"#{RED}")
    fig.text(.5,.018,"Metrik tidak berasal dari head-to-head terkontrol; jangan simpulkan kausalitas arsitektur/DAM.",ha="center",fontsize=8,color=f"#{RED}")
    fig.tight_layout(rect=[0,.13,1,.95])
    fig.savefig(ASSET_DIR/"perbandingan_metrik.png",dpi=200,bbox_inches="tight")
    plt.close(fig)


def data_figure() -> None:
    fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.7))
    ax=axes[0]
    y=[1,0]
    ax.barh(1,5374,color=f"#{TEAL}",label="gambar berlabel pasangan")
    ax.barh(0,2419,color=f"#{TEAL}")
    ax.barh(0,4309,left=2419,color=f"#{RED}",label="gambar tanpa label pasangan")
    ax.set_yticks(y,["Train","Validation"]); ax.set_title("Fish4Knowledge",weight="bold")
    ax.set_xlabel("Jumlah gambar"); ax.legend(fontsize=7,loc="lower right")
    ax.text(500,1,"5.374 paired",va="center",color="white",weight="bold",fontsize=8)
    ax.text(500,0,"2.419 paired",va="center",color="white",weight="bold",fontsize=8)
    ax.text(4550,0,"4.309\nno label",va="center",ha="center",color="white",weight="bold",fontsize=8)
    ax.text(.5,-.24,"+ 4.316 orphan labels (train); +3 (val)",transform=ax.transAxes,ha="center",fontsize=8,color=f"#{RED}")

    ax=axes[1]
    ax.barh([1,0],[929,226],color=f"#{TEAL}",label="image-label pairs")
    ax.set_yticks([1,0],["Train","Validation"]); ax.set_title("FishDisease · lesion",weight="bold")
    ax.set_xlabel("Jumlah pasangan gambar-label"); ax.text(480,1,"1.705 boxes",ha="center",va="center",color="white",weight="bold")
    ax.text(115,0,"441 boxes",ha="center",va="center",color="white",weight="bold")

    ax=axes[2]
    ax.bar([0,1],[250,100],color=[f"#{NAVY}",f"#{GOLD}"],width=.6)
    ax.set_xticks([0,1],["Train","Validation"]); ax.set_title("Kaggle classification",weight="bold")
    ax.set_ylabel("Gambar per kelas (masing-masing 7 kelas)"); ax.set_ylim(0,300)
    for x,val in enumerate([250,100]): ax.text(x,val+7,str(val),ha="center",weight="bold")
    ax.text(.5,-.24,"Laporan Tugas 2: seluruh validation overlap dengan train",transform=ax.transAxes,ha="center",fontsize=7.6,color=f"#{RED}")
    for ax in axes:
        ax.grid(axis="x",alpha=.4); ax.set_axisbelow(True)
    fig.suptitle("Ukuran dataset dan temuan kualitas data",fontsize=14,weight="bold",color=f"#{NAVY}")
    fig.tight_layout(rect=[0,.10,1,.92])
    fig.savefig(ASSET_DIR/"audit_dataset.png",dpi=200,bbox_inches="tight")
    plt.close(fig)


def training_curve_figure() -> None:
    specs=[
        ("Deteksi ikan","runs/detect/Runs_DynamicAttention/model_ikan-4/results.csv",[
            ("metrics/precision(B)","Precision"),("metrics/recall(B)","Recall"),("metrics/mAP50(B)","mAP50"),("metrics/mAP50-95(B)","mAP50–95")]),
        ("Deteksi lesi","runs/detect/Runs_DynamicAttention/model_lesi/results.csv",[
            ("metrics/precision(B)","Precision"),("metrics/recall(B)","Recall"),("metrics/mAP50(B)","mAP50"),("metrics/mAP50-95(B)","mAP50–95")]),
        ("Klasifikasi","runs/classify/Runs_DynamicAttention/model_penyakit/results.csv",[
            ("metrics/accuracy_top1","Top-1"),("metrics/accuracy_top5","Top-5")]),
    ]
    fig,axes=plt.subplots(1,3,figsize=(13.5,4.1))
    palette=[NAVY,RED,TEAL,GOLD]
    for ax,(title,path,series) in zip(axes,specs):
        frame=pd.read_csv(ROOT/path); frame.columns=[c.strip() for c in frame.columns]
        for i,(col,label) in enumerate(series):
            ax.plot(frame["epoch"],frame[col],marker="o",lw=2,color=f"#{palette[i]}",label=label)
        ax.set_title(title,weight="bold"); ax.set_xlabel("Epoch"); ax.set_xticks(frame["epoch"])
        ax.set_ylim(0,1.04); ax.grid(alpha=.7); ax.legend(fontsize=8,frameon=False)
    fig.suptitle("Validation per epoch · YOLO26 + DynamicAttention",fontsize=14,weight="bold",color=f"#{NAVY}")
    fig.text(.5,-.015,"Satu run/seed; kurva tidak mengatasi kebocoran split atau ketidakcocokan anotasi.",ha="center",fontsize=8,color=f"#{RED}")
    fig.tight_layout(rect=[0,.04,1,.91])
    fig.savefig(ASSET_DIR/"kurva_training_dynamic_attention.png",dpi=200,bbox_inches="tight")
    plt.close(fig)


def shade_cell(cell, color: str) -> None:
    tc_pr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),color); tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=8.5) -> None:
    cell.text=""
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
    run=p.add_run(text); run.bold=bold; run.font.name="Aptos"; run.font.size=Pt(size)
    if color: run.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run("YOLOComVis · Tugas 4 · Halaman ")
    run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GRAY)
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); paragraph._p.append(field)


def inline_format(paragraph, text: str) -> None:
    text=re.sub(r"\[([^\]]+)\]\(([^)]+)\)",r"\1",text)
    parts=re.split(r"(\*\*.*?\*\*|`.*?`|\*[^*]+\*)",text)
    for part in parts:
        if not part: continue
        bold=part.startswith("**") and part.endswith("**")
        italic=part.startswith("*") and part.endswith("*") and not bold
        code=part.startswith("`") and part.endswith("`")
        value=part[2:-2] if bold else part[1:-1] if italic or code else part
        run=paragraph.add_run(value); run.bold=bold; run.italic=italic
        if code:
            run.font.name="Consolas"; run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(TEAL)


def render_docx() -> None:
    doc=Document()
    section=doc.sections[0]
    section.top_margin=Inches(.72); section.bottom_margin=Inches(.68)
    section.left_margin=Inches(.78); section.right_margin=Inches(.78)
    styles=doc.styles
    styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(9.2)
    styles["Normal"].paragraph_format.space_after=Pt(5)
    for name,size in [("Title",25),("Heading 1",17),("Heading 2",12.5),("Heading 3",10.5)]:
        st=styles[name]; st.font.name="Aptos Display"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY)
        st.paragraph_format.keep_with_next=True
    header=section.header.paragraphs[0]
    header.text="UNIVERSITAS BRAWIJAYA  |  COMPUTER VISION  |  KELOMPOK 15"
    header.style="Caption"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(section.footer.paragraphs[0])

    lines=MD_PATH.read_text(encoding="utf-8").splitlines()
    i=0; first_title=True; cover=True
    while i<len(lines):
        raw=lines[i]; line=raw.strip()
        if not line or line=="---": i+=1; continue
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title" if first_title else "Heading 1")
            inline_format(p,line[2:]);
            if first_title:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(70); p.paragraph_format.space_after=Pt(22)
                first_title=False
            i+=1; continue
        if line.startswith("## "):
            title=line[3:]
            if title=="Ringkasan eksekutif" and cover:
                doc.add_page_break(); cover=False
            p=doc.add_paragraph(style="Heading 2")
            inline_format(p,title)
            if cover: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            i+=1; continue
        if line.startswith("### "):
            p=doc.add_paragraph(style="Heading 3"); inline_format(p,line[4:]); i+=1; continue
        if line.startswith("!") and i+1>0:
            match=re.match(r"!\[(.*?)\]\((.*?)\)",line)
            if match:
                caption,path=match.groups(); image_path=(REPORT_DIR/path).resolve()
                if image_path.exists():
                    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(image_path),width=Inches(6.55))
                    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    r=cap.add_run(caption); r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
                i+=1; continue
        if line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                cells=[c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?",c or "-") for c in cells): rows.append(cells)
                i+=1
            if rows:
                ncols=max(len(r) for r in rows); table=doc.add_table(rows=1,cols=ncols); table.style="Light Shading Accent 1"
                table.autofit=True
                for ri,row in enumerate(rows):
                    cells=table.rows[0].cells if ri==0 else table.add_row().cells
                    if len(cells)<ncols: continue
                    for ci in range(ncols):
                        value=row[ci] if ci<len(row) else ""
                        value=re.sub(r"\*\*(.*?)\*\*",r"\1",value); value=re.sub(r"`([^`]*)`",r"\1",value)
                        set_cell_text(cells[ci],value,bold=(ri==0),color="FFFFFF" if ri==0 else None,size=7.6)
                        if ri==0: shade_cell(cells[ci],NAVY)
                    tr_pr=table.rows[-1]._tr.get_or_add_trPr(); tr_pr.append(OxmlElement("w:cantSplit"))
                doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        if line.startswith("> "):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.right_indent=Inches(.18)
            inline_format(p,line[2:]);
            for r in p.runs: r.italic=True; r.font.color.rgb=RGBColor.from_string(TEAL)
            i+=1; continue
        if line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); inline_format(p,line[2:]); i+=1; continue
        if re.match(r"\d+\. ",line):
            p=doc.add_paragraph(style="List Number"); inline_format(p,re.sub(r"^\d+\. ","",line)); i+=1; continue
        paragraph_lines=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |\| |\> |\- |\d+\. |!)",lines[i].strip()):
            paragraph_lines.append(lines[i].strip()); i+=1
        p=doc.add_paragraph()
        if cover:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(12)
            for s in p.runs: s.font.color.rgb=RGBColor.from_string(GRAY)
        inline_format(p," ".join(paragraph_lines))
        if cover and len(paragraph_lines)==1:
            for run in p.runs: run.font.size=Pt(10); run.font.color.rgb=RGBColor.from_string(GRAY)
    doc.core_properties.title="Tugas 4 — Laporan Akhir YOLOComVis"
    doc.core_properties.subject="Evaluasi perkembangan, arsitektur, dan kelemahan model"
    doc.core_properties.author="Kelompok 15 — Teknik Komputer, Universitas Brawijaya"
    doc.save(DOCX_PATH)


def main() -> None:
    ASSET_DIR.mkdir(parents=True,exist_ok=True)
    theme(); architecture_figure(); metric_figure(); data_figure(); training_curve_figure(); render_docx()
    print(f"Markdown: {MD_PATH}")
    print(f"DOCX: {DOCX_PATH}")
    print(f"Figures: {ASSET_DIR}")


if __name__=="__main__":
    main()
